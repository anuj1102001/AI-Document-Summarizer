import io

import docx
import pytest

from app.errors import AppError
from app.ingest.chunking import chunk_pages, looks_like_heading
from app.ingest.parsers import Page, clean_text, detect_kind, parse
from tests.conftest import make_image_only_pdf, make_pdf


def test_pdf_pages_and_empty_page_detection():
    parsed = parse(make_pdf(["First page about refunds.", "", "Third page about renewals."]), "a.pdf", 50)
    assert parsed.kind == "pdf"
    assert [p.number for p in parsed.pages] == [1, 2, 3]
    assert "refunds" in parsed.pages[0].text
    assert parsed.empty_pages == [2]


def test_image_only_pdf_is_reported_as_scanned():
    with pytest.raises(AppError) as e:
        parse(make_image_only_pdf(), "scan.pdf", 50)
    assert e.value.code == "NO_TEXT" and "OCR" in e.value.message


def test_encrypted_pdf_is_rejected_clearly():
    with pytest.raises(AppError) as e:
        parse(make_pdf(["secret"], password="pw"), "locked.pdf", 50)
    assert e.value.code == "ENCRYPTED"


def test_page_limit():
    with pytest.raises(AppError) as e:
        parse(make_pdf(["x" * 40] * 5), "long.pdf", 3)
    assert e.value.code == "TOO_MANY_PAGES"


def test_type_is_detected_from_bytes_not_name():
    with pytest.raises(AppError) as e:
        detect_kind(b"MZ\x90\x00 not a pdf", "invoice.pdf")
    assert e.value.code == "UNSUPPORTED_TYPE"
    assert detect_kind(make_pdf(["hello world text"]), "renamed.txt") == "pdf"


def test_docx_paragraphs_and_tables():
    d = docx.Document()
    d.add_heading("Payment Terms", 1)
    d.add_paragraph("Invoices are payable within 30 days of receipt.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Late fee"
    t.rows[0].cells[1].text = "1.5% per month"
    buf = io.BytesIO()
    d.save(buf)
    parsed = parse(buf.getvalue(), "terms.docx", 50)
    assert parsed.kind == "docx"
    text = parsed.pages[0].text
    assert "payable within 30 days" in text and "Late fee | 1.5% per month" in text


def test_text_must_be_utf8():
    with pytest.raises(AppError) as e:
        parse("café".encode("latin-1"), "notes.txt", 50)
    assert e.value.code == "BAD_ENCODING"


def test_clean_text_rejoins_hyphenated_words_and_keeps_paragraphs():
    assert clean_text("infor-\nmation is   here\n\n\n\nnext") == "information is here\n\nnext"


@pytest.mark.parametrize(
    "line,expected",
    [
        ("2.1 Scope of Services", True),
        ("TERMINATION", True),
        ("Payment Terms", True),
        ("The supplier shall deliver the goods.", False),
        ("What is covered?", False),
    ],
)
def test_heading_heuristic(line, expected):
    assert looks_like_heading(line) is expected


def test_chunks_keep_pages_headings_and_size():
    body = " ".join(f"Sentence {i} explains a detail of the agreement." for i in range(60))
    pages = [Page(1, f"1. Definitions\n\n{body}"), Page(2, f"2. Payment Terms\n\n{body}")]
    chunks = chunk_pages(pages, size=600, overlap=80)
    assert all(len(c.text) <= 700 for c in chunks)
    assert chunks[0].page_start == 1 and chunks[-1].page_end == 2
    assert {c.heading for c in chunks} >= {"1. Definitions", "2. Payment Terms"}
    assert all(c.page_end - c.page_start <= 1 for c in chunks)
    # consecutive chunks overlap slightly so a fact split at a boundary is still retrievable
    assert any(chunks[i].text[:30] in chunks[i - 1].text for i in range(1, len(chunks)))


def test_single_short_document_is_one_chunk():
    chunks = chunk_pages([Page(1, "A short note about the meeting on Friday.")])
    assert len(chunks) == 1 and chunks[0].page_start == 1


def test_titles_prefer_metadata_then_meaningful_filenames_then_first_heading():
    from app.services.documents import choose_title

    assert choose_title("Annual Report 2025", "x.pdf", "") == "Annual Report 2025"
    assert choose_title(None, "Q3_board-pack.pdf", "Cover") == "Q3 board pack"
    assert choose_title(None, "scan_0042.pdf", "Lease Agreement for Flat 4B\nThis lease...") == "Lease Agreement for Flat 4B"
    assert choose_title(None, "msa.pdf", "Master Services Agreement\n1. Parties") == "Master Services Agreement"
    assert choose_title(None, "doc1.pdf", "This first line is a sentence.") == "doc1"
